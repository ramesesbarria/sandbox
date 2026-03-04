from docx.table import Table
from copy import deepcopy
from utils.helper import iter_block_items, add_cell_to_header, merge_cells_range, align_cell_text
from utils.helper import multi_table_block_finder_with_tag, set_table_full_width

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

TABLE_IDENTIFIERS = {
    tuple(['Investment', 'Amount', 'Sell cost']): {
		"name":"Fix Issues encounter in merge multiple split table",
        "process": 4,
		"processDescription": "Fix Issues encounter in merge multiple split",
	},
	tuple(['Investment', 'Amount', 'Buy cost']): {
		"name":"Fix Issues encounter in merge multiple split table",
        "process": 4,
		"processDescription": "Fix Issues encounter in merge multiple split",
	},
	tuple(['Investment', 'Amount', 'Ongoing Investment\nManagement Fees']): {
		"name":"Fix Issues encounter in merge multiple split table",
        "process": 4,
		"processDescription": "Fix Issues encounter in merge multiple split",
	},
}


def post_fix_table_cleanup(document):
    """Post-processing of tables after fixing fragmented tables.
    Adds extra header columns, merges cells, aligns text, and sets column widths.
    """
    datas = multi_table_block_finder_with_tag(document, TABLE_IDENTIFIERS)
   
    for data in datas:
        table = data[0]
        add_cell_to_header(table.rows[0], new_header_text="Extra Column")
        merge_cells_range(table.rows[0],2,3,10)
        
        if data[1] == "Ongoing Investment\nManagement Fees":
            align_cell_text(table.rows[0], -1, alignment=WD_ALIGN_PARAGRAPH.RIGHT)

            add_cell_to_header(table.rows[-1], new_header_text="Extra Column")
            merge_cells_range(table.rows[-1],2,3,10)
            align_cell_text(table.rows[-1], -1, alignment=WD_ALIGN_PARAGRAPH.RIGHT)

            add_cell_to_header(table.rows[-2], new_header_text="Extra Column")
            merge_cells_range(table.rows[-2],2,3,10)
            align_cell_text(table.rows[-2], -1, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            
        elif data[1]  == "Sell cost" or data[1] == "Buy cost":
            align_cell_text(table.rows[0], -1, alignment=WD_ALIGN_PARAGRAPH.CENTER)

        for row in table.rows:
            if len(row.cells) > 3:
                row.cells[3].width = Inches(1)


def fix_fragmented_tables(doc, merge_condition=None):

    block_list = list(iter_block_items(doc))

    i = 0

    tables_to_remove = []
    
    while i < len(block_list):
        block = block_list[i]

        if isinstance(block, Table):


            # Keep merging following tables as long as they are consecutive
            j = i + 1
            while j < len(block_list) and isinstance(block_list[j], Table):
                next_table = block_list[j]
                # Append all rows from next_table to current table
                for row in next_table.rows:
                    new_row = deepcopy(row._tr)
                    block._tbl.append(new_row)
                tables_to_remove.append(next_table)
                j += 1
            i = j  # Skip merged tables
        else:
            i += 1

    # Remove merged tables from the document
    for t in tables_to_remove:
        t._element.getparent().remove(t._element)

    # Fix table issues after merging split tables.
    post_fix_table_cleanup(doc)

    return doc